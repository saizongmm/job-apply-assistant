# -*- coding: utf-8 -*-
"""Build a tailored resume .docx from a parsed profile + a tailored JSON.

NO hardcoded candidate data. Everything comes from:
  --profile  profile.json   (from parse_cv.py)
  --tailored tailored.json  (LLM-produced: target + summary bullets + optional experience)

tailored.json shape:
{
  "target": {"title": "...", "company": "...", "location": "..."},
  "name": "override?" , "tagline": "...", "email": "...", "phone": "...",
  "summary": ["bullet1", "bullet2", ...],          # REQUIRED (the tailoring)
  "experience": [ {"company":..,"role":..,"period":..,"bullets":[..]} ],  # optional override
  "education": [ ["school","degree"], ... ]          # optional override
}

Usage:
  python3 make_docx.py --profile profile.json --tailored tailored.json --out resume.docx
"""
import sys, json, os
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Calibri"


def set_run(r, size, bold, color, italic=False, name=FONT):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = name
    r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:eastAsia'), name)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    set_run(r, 12, True, NAVY)
    pPr = p._element.get_or_add_pPr()
    pbdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pbdr.makeelement(qn('w:bottom'),
                              {qn('w:val'): 'single', qn('w:sz'): '6',
                               qn('w:space'): '1', qn('w:color'): '1F4E79'})
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def build(profile, tailored):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    name = tailored.get("name") or profile.get("name") or "CANDIDATE NAME"
    tagline = tailored.get("tagline") or profile.get("tagline") or ""
    email = tailored.get("email") or profile.get("email") or ""
    phone = tailored.get("phone") or profile.get("phone") or ""

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(name), 22, True, NAVY)

    if tagline:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(tagline), 10.5, False, GRAY)

    contact = "  |  ".join([c for c in (email, phone) if c])
    if contact:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        set_run(p.add_run(contact), 10, False, GRAY)

    tgt = tailored.get("target") or {}
    tline = ""
    if tgt.get("title"):
        tline += f"Target: {tgt['title']}"
    if tgt.get("company"):
        tline += f" - {tgt['company']}"
    if tgt.get("location"):
        tline += f" ({tgt['location']})"
    if tline:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
        set_run(p.add_run(tline), 10.5, True, NAVY)

    # Summary (the tailored part)
    summary = tailored.get("summary") or profile.get("summary") or []
    if summary:
        add_heading(doc, "Professional Summary")
        for s in summary:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.15
            set_run(p.add_run(s), 10.5, False, BLACK)

    # Experience
    exp = tailored.get("experience") or profile.get("experience") or []
    if exp:
        add_heading(doc, "Professional Experience")
        for e in exp:
            head = e.get("company") or ""
            role = e.get("role") or ""
            period = e.get("period") or ""
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(head), 11, True, BLACK)
            if role:
                set_run(p.add_run(f"   {role}"), 10.5, True, NAVY)
            if period:
                p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(2)
                set_run(p2.add_run(period), 9.5, False, GRAY, italic=True)
            for b in e.get("bullets", []):
                pb = doc.add_paragraph(style='List Bullet')
                pb.paragraph_format.space_after = Pt(1); pb.paragraph_format.line_spacing = 1.1
                set_run(pb.add_run(b), 10, False, BLACK)

    # Education
    edu = tailored.get("education") or profile.get("education") or []
    if edu:
        add_heading(doc, "Education")
        for item in edu:
            if isinstance(item, (list, tuple)) and len(item) >= 2:
                school, degree = item[0], item[1]
            elif isinstance(item, dict):
                school = item.get("school") or item.get("company") or ""
                degree = item.get("degree") or ""
            else:
                continue
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
            set_run(p.add_run(f"{school} - {degree}"), 10.5, False, BLACK)

    return doc


def main():
    ap = __import__("argparse").ArgumentParser()
    ap.add_argument("--profile", required=True)
    ap.add_argument("--tailored", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    profile = json.loads(Path(args.profile).read_text(encoding="utf-8"))
    tailored = json.loads(Path(args.tailored).read_text(encoding="utf-8"))

    doc = build(profile, tailored)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(f"已生成定制简历: {args.out}")


if __name__ == "__main__":
    main()
