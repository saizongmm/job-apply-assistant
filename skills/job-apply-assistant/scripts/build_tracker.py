# -*- coding: utf-8 -*-
"""Append a harvested jobs TSV into the application tracker (dedup + normalize).

Reusable version of the Selina pipeline's append logic. No hardcoded candidate data.

TSV format (from WebFetch compact prompt): one row per job, 5 columns, TAB-separated:
   JOB_TITLE <tab> COMPANY <tab> LOCATION <tab> DATE_POSTED <tab> URL
The URL must be the full linkedin.com/jobs/view/NNNNN link (job-id is extracted from it).
If a row's 5th col is a bare numeric id instead of a URL, it is accepted too.

Usage:
  python3 build_tracker.py --tsv harvest.tsv --tracker application_tracker.xlsx
  python3 build_tracker.py --tsv harvest.tsv --tracker application_tracker.xlsx --create
"""
import sys, re, shutil, os
from pathlib import Path
import openpyxl

HEADER = ["#", "City", "Company", "Job Title", "Seniority", "Status",
          "LinkedIn URL", "Tailored Resume", "Cover Letter", "Applied?", "Reply", "Notes"]

CITY_MAP = {
    "shanghai": "Shanghai", "beijing": "Beijing", "shenzhen": "Shenzhen",
    "guangzhou": "Guangzhou", "hangzhou": "Hangzhou", "chengdu": "Chengdu",
    "wuhan": "Wuhan", "nanjing": "Nanjing", "suzhou": "Suzhou", "tianjin": "Tianjin",
    "xian": "Xi'an", "xi'an": "Xi'an", "ningbo": "Ningbo", "hefei": "Hefei",
    "zhengzhou": "Zhengzhou", "changsha": "Changsha", "baoding": "Baoding",
    "dongguan": "Dongguan", "jinan": "Jinan", "kunshan": "Kunshan",
    "hong kong": "Hong Kong", "hongkong": "Hong Kong", "sg": "Singapore",
    "singapore": "Singapore", "london": "London", "new york": "New York",
    "san francisco": "San Francisco", "sydney": "Sydney", "tokyo": "Tokyo",
    "dublin": "Dublin", "berlin": "Berlin", "amsterdam": "Amsterdam",
}


def norm_city(loc, default="Unspecified"):
    if not loc:
        return default
    low = loc.lower()
    if "china" not in low and "shanghai" not in low and "hong kong" not in low and "beijing" not in low:
        # could be a foreign country string; keep it
        pass
    m = re.search(r'greater\s+([a-z]+)\s+area', low)
    if m and m.group(1) in CITY_MAP:
        return CITY_MAP[m.group(1)]
    m = re.search(r'([a-z\'\u4e00-\u9fff]+)\s*-\s*[a-z\u4e00-\u9fff]+\s+metropolitan', low)
    if m and m.group(1) in CITY_MAP:
        return CITY_MAP[m.group(1)]
    m = re.search(r'([a-z\'\u4e00-\u9fff]+)\s+metropolitan', low)
    if m and m.group(1) in CITY_MAP:
        return CITY_MAP[m.group(1)]
    first = loc.split(",")[0].strip().lower()
    if first in CITY_MAP:
        return CITY_MAP[first]
    for k, v in CITY_MAP.items():
        if k in low:
            return v
    # fallback: title-case the primary token
    return loc.split(",")[0].strip().title() or default


def infer_role(title):
    t = (title or "").lower()
    if re.search(r'\b(director|cmo|cto|cro|cfo|vp|head|avp|svp|chief|lead)\b', t):
        return "Director"
    if re.search(r'communication|pr\b|marcom|comms|公关|传播|企业传播|公共事务|媒介', t):
        return "Comms"
    if re.search(r'\b(intern|实习)\b', t):
        return "Intern"
    if re.search(r'\b(junior|entry|graduate|associate|analyst)\b', t) and "senior" not in t:
        return "IC"
    if re.search(r'\b(senior|sr\.?|staff|principal|专家)\b', t):
        return "Senior"
    return "Manager"


def job_id_from(url):
    if not url:
        return None
    m = re.search(r'/jobs/view/.*?(\d+)(?:[?#]|\s|$)', str(url))
    if m:
        return m.group(1)
    m = re.search(r'(\d{6,})', str(url))
    if m:
        return m.group(1)
    return None


def score_row(title, company, city, role, wanted_kw, wanted_cities):
    t = (str(title) + " " + str(company)).lower()
    score = 0
    for kw in wanted_kw:
        if kw.lower() in t:
            score += 2
    if any(c.lower() in str(city).lower() for c in wanted_cities):
        score += 1
    if role in ("Director", "Comms", "Senior"):
        score += 1
    return score


def main():
    ap = __import__("argparse").ArgumentParser()
    ap.add_argument("--tsv", required=True)
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--create", action="store_true", help="tracker 不存在则新建(带表头)")
    ap.add_argument("--wanted", default="", help="目标岗位关键词(逗号分隔,用于评分)")
    ap.add_argument("--cities", default="", help="目标城市(逗号分隔,用于评分)")
    args = ap.parse_args()

    tsv = Path(args.tsv)
    tracker = Path(args.tracker)
    if not tsv.exists():
        print(f"找不到 TSV: {tsv}")
        sys.exit(1)

    if tracker.exists():
        shutil.copy(tracker, tracker.with_name(tracker.stem + "_backup_before_append.xlsx"))
        wb = openpyxl.load_workbook(tracker)
        ws = wb.active
    elif args.create:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADER)
    else:
        print(f"tracker 不存在: {tracker} (加 --create 新建)")
        sys.exit(1)

    existing_ids = set()
    for r in range(2, ws.max_row + 1):
        url = ws.cell(r, 7).value
        jid = job_id_from(url)
        if jid:
            existing_ids.add(jid)

    wanted = [k.strip() for k in args.wanted.split(",") if k.strip()]
    wanted_cities = [c.strip() for c in args.cities.split(",") if c.strip()]

    rows = []
    with open(tsv, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\n")
            if not line.strip():
                continue
            parts = line.split("\t")
            if len(parts) < 5:
                continue
            title, company, location, date, url = (parts[0], parts[1], parts[2], parts[3], parts[4])
            rows.append((title, company, location, date, url))

    seen = set()
    new = dups_existing = dups_within = 0
    cur = 0
    for r in range(2, ws.max_row + 1):
        try:
            cur = max(cur, int(ws.cell(r, 1).value))
        except (TypeError, ValueError):
            pass

    for title, company, location, date, url in rows:
        jid = job_id_from(url)
        if not jid:
            continue
        if jid in existing_ids:
            dups_existing += 1
            continue
        if jid in seen:
            dups_within += 1
            continue
        seen.add(jid)
        cur += 1
        city = norm_city(location)
        role = infer_role(title)
        full_url = url if str(url).startswith("http") else "https://www.linkedin.com/jobs/view/" + jid
        notes = f"harvest {date}; score={score_row(title, company, city, role, wanted, wanted_cities)}"
        ws.append([cur, city, company, title, role, "LIVE", full_url, "", "", "", "", notes])
        new += 1

    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    wb.save(tracker)
    print(f"dups vs existing: {dups_existing} | dups within batch: {dups_within} | NEW appended: {new}")
    print(f"new max #: {cur} | new max_row: {ws.max_row}")


if __name__ == "__main__":
    main()
