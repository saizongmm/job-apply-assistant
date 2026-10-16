#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Path A opener — generic version (no hardcoded candidate/tracker).

Opens each job's URL in the USER's already-logged-in Chrome via `open -a "Google Chrome"`,
so there is no cookie-portability or anti-bot issue. The user clicks Apply manually;
then reports back which # they applied, and we mark Applied?=Yes.

Filters (honors "don't re-open things already handled"):
  - skip Applied? in {yes, y, true, no, n, declined, skip, opened, open}
  - only open rows that HAVE a Tailored Resume file on disk

Usage:
  python3 open_jobs.py --tracker application_tracker.xlsx --dry-run
  python3 open_jobs.py --tracker application_tracker.xlsx --auto-open --batch 15 --delay 2
  python3 open_jobs.py --tracker application_tracker.xlsx --only 252
  python3 open_jobs.py --tracker application_tracker.xlsx --done 252 299
  python3 open_jobs.py --tracker application_tracker.xlsx --skip 42 46
"""
import os, sys, subprocess, datetime, argparse
from pathlib import Path

SKIP_APPLIED = ("yes", "y", "true", "已投", "applied", "no", "n", "declined",
                "skip", "false", "opened", "open", "打开")
REQUIRE_RESUME = True
_WS = None
_COLS = None
_JOBS = None


def load_jobs(tracker):
    import openpyxl
    wb = openpyxl.load_workbook(tracker)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    cols = {h: i for i, h in enumerate(hdr)}
    resume_dir = tracker.parent / "tailored_resumes"
    jobs = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for h, i in cols.items()}
        row["_rownum"] = r
        fn = row.get("Tailored Resume")
        rp = (resume_dir / str(fn)) if fn else None
        row["_resume_path"] = rp
        row["_has_resume"] = bool(fn) and rp is not None and rp.exists()
        jobs.append(row)
    return wb, ws, cols, jobs


def open_in_chrome(url):
    """Open a URL in the user's installed Chrome if possible, else the default
    browser. Cross-platform (macOS / Linux / Windows). The browser must already
    be logged in to the job site — we never pass cookies or credentials."""
    import sys
    plat = sys.platform
    if plat == "darwin":
        for cmd in (["open", "-a", "Google Chrome", url], ["open", url]):
            try:
                subprocess.run(cmd, check=True)
                return True
            except Exception:
                continue
        return False
    if plat.startswith("win"):
        candidates = [
            [r"C:\Program Files\Google\Chrome\Application\chrome.exe", url],
            [r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe", url],
            ["cmd", "/c", "start", "", url],
        ]
        for cmd in candidates:
            try:
                subprocess.run(cmd, check=True, shell=False)
                return True
            except Exception:
                continue
        return False
    # Linux / other POSIX
    for cmd in (["google-chrome", url], ["chromium", url], ["chromium-browser", url],
                ["xdg-open", url]):
        try:
            subprocess.run(cmd, check=True)
            return True
        except Exception:
            continue
    return False


def ts():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M")


def row_by_num(jobs, num):
    for j in jobs:
        if j.get("#") == num:
            return j
    return None


def mark_opened(job):
    r = job["_rownum"]
    _WS.cell(r, _COLS["Applied?"] + 1).value = "Opened"
    rp = job.get("_resume_path")
    note = f"{ts()} PathA opened; resume={rp.name if rp else '?'}"
    _WS.cell(r, _COLS["Notes"] + 1).value = note
    _WS.parent.save(_TRACKER)


def mark_done(nums):
    n = 0
    for num in nums:
        j = row_by_num(_JOBS, num)
        if not j:
            print(f"  ! 找不到 #{num}")
            continue
        r = j["_rownum"]
        _WS.cell(r, _COLS["Applied?"] + 1).value = "Yes"
        _WS.cell(r, _COLS["Notes"] + 1).value = f"{ts()} PathA applied"
        n += 1
    _WS.parent.save(_TRACKER)
    print(f"  已标记 {n} 条为 Applied?=Yes。")


def mark_skip(nums):
    n = 0
    for num in nums:
        j = row_by_num(_JOBS, num)
        if not j:
            print(f"  ! 找不到 #{num}")
            continue
        r = j["_rownum"]
        _WS.cell(r, _COLS["Applied?"] + 1).value = "No"
        _WS.cell(r, _COLS["Notes"] + 1).value = f"{ts()} PathA user declined"
        n += 1
    _WS.parent.save(_TRACKER)
    print(f"  已标记 {n} 条为 Applied?=No。")


def build_target(jobs, start_from=None, only=None):
    out = []
    for j in jobs:
        if only is not None:
            if j.get("#") == only:
                out.append(j)
            continue
        if start_from is not None and j.get("#") < start_from:
            continue
        if str(j.get("Applied?")).strip().lower() in SKIP_APPLIED:
            continue
        if REQUIRE_RESUME and not j.get("_has_resume"):
            continue
        out.append(j)
    return out


def run_auto_open(target, batch=15, delay=2):
    print(f"\n=== Path A auto-open: {len(target)} 条, 本批上限 {batch}, 间隔 {delay}s ===")
    opened = failed = n = 0
    for j in target:
        if n >= batch:
            break
        url = j.get("LinkedIn URL")
        rp = j.get("_resume_path")
        print(f"  [{opened + failed + 1}/{len(target)}] 打开 #{j.get('#')} {j.get('Company')} — {j.get('Job Title')}")
        if not url:
            print("     !! 无 URL,跳过")
            continue
        if open_in_chrome(url):
            mark_opened(j)
            opened += 1
            n += 1
            if rp:
                print(f"     简历: {rp}")
        else:
            failed += 1
            print("     !! 打开失败(确认 Chrome 已安装且在本机 GUI 会话)")
        if n < batch:
            import time
            time.sleep(delay)
    print(f"\n本批已打开 {opened}, 失败 {failed}。其余待下一批(已打开不重复)。")


def main():
    global _WS, _COLS, _JOBS, _TRACKER
    ap = argparse.ArgumentParser()
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--auto-open", action="store_true")
    ap.add_argument("--batch", type=int, default=15)
    ap.add_argument("--delay", type=float, default=2.0)
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--from", dest="start_from", type=int, default=None)
    ap.add_argument("--only", type=int, default=None)
    ap.add_argument("--done", nargs="*", type=int, default=None)
    ap.add_argument("--skip", nargs="*", type=int, default=None)
    args = ap.parse_args()

    _TRACKER = Path(args.tracker).resolve()
    if not _TRACKER.exists():
        print(f"找不到 tracker: {_TRACKER}")
        sys.exit(1)

    _wb, _WS, _COLS, _JOBS = load_jobs(_TRACKER)

    if args.done is not None:
        mark_done(args.done)
        return
    if args.skip is not None:
        mark_skip(args.skip)
        return

    target = build_target(_JOBS, start_from=args.start_from, only=args.only)
    if args.dry_run:
        print(f"计划打开 {len(target)} 条(有简历且未处理)")
        for j in target:
            print(f"  #{j.get('#')} [{j.get('City')}] {j.get('Company')} — {j.get('Job Title')}")
        return
    if not target:
        print("没有需要打开的条目。")
        return
    if args.auto_open:
        run_auto_open(target, batch=args.batch, delay=args.delay)
        return
    if args.all:
        ok = 0
        for j in target:
            if open_in_chrome(j.get("LinkedIn URL")):
                ok += 1
        print(f"已请求打开 {ok}/{len(target)}。完成后运行 --done <#列表>。")
        return
    # default: just open the first one as a smoke test
    if target:
        open_in_chrome(target[0].get("LinkedIn URL"))
        mark_opened(target[0])


if __name__ == "__main__":
    main()
